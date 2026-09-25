import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color of a cell using XML shading."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding (in dxa: 20 dxa = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_doc(output_path="PRJ_409_Implementation_Stack.docx"):
    doc = Document()

    # Set page margins to 0.75 in (54 pt)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Document Header / Title
    title = doc.add_paragraph()
    title_run = title.add_run("PRJ_409: AI Personal Finance Advisor")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    title.paragraph_format.space_after = Pt(2)

    subtitle = doc.add_paragraph()
    sub_run = subtitle.add_run("Technical Specification & System Implementation Stack")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(12)
    sub_run.font.color.rgb = RGBColor(100, 116, 139) # Slate 500
    subtitle.paragraph_format.space_after = Pt(16)

    # Table Header Data
    headers = [
        "Architectural Layer",
        "Component / Tool",
        "Version",
        "Technical Role & Purpose in PRJ_409"
    ]

    data = [
        ["Frontend Framework", "React", "19.2", "Component-based Single Page Application (SPA) architecture"],
        ["Programming Language", "TypeScript", "6.0", "Static typing, interface contracts, and compile-time error prevention"],
        ["Build Tool & Bundler", "Vite", "8.2", "High-speed Hot Module Replacement (HMR) and optimized asset bundling"],
        ["Styling & Design", "TailwindCSS", "4.3", "Utility-first responsive design, modern dark/light UI tokens"],
        ["Data Visualization", "Recharts", "3.10", "Interactive SVG financial trajectory charts (Bar charts & Pie charts)"],
        ["Iconography", "Lucide React", "1.35", "Modern financial and dashboard iconography"],
        ["Client Routing", "React Router DOM", "7.18", "Declarative client-side routing with protected auth guards"],
        ["HTTP Client", "Axios", "1.20", "Centralized API client with JWT request/response interceptors"],
        ["Backend Framework", "FastAPI", "0.115", "Asynchronous high-performance RESTful API micro-framework"],
        ["Backend Language", "Python", "3.10", "Core computational language for business logic, math, and ML"],
        ["ASGI Server", "Uvicorn", "0.32", "Production-grade asynchronous web server gateway"],
        ["Database ORM", "SQLAlchemy", "2.0", "Object Relational Mapping (ORM) and relational schema management"],
        ["Data Validation", "Pydantic v2", "2.0+", "Automatic schema serialization, request validation, and typing"],
        ["Authentication & Tokens", "Python-Jose", "3.3", "JSON Web Token (JWT) encoding/decoding using HMAC-SHA256"],
        ["Password Security", "Passlib & Bcrypt", "4.0", "Cryptographic password hashing with automatic salt generation"],
        ["Machine Learning", "Scikit-Learn", "1.3", "TF-IDF Vectorization, Logistic Regression, Evaluation metrics"],
        ["Data Manipulation", "Pandas & NumPy", "2.0+", "Matrix operations, dataset preprocessing, and financial aggregations"],
        ["Model Persistence", "Joblib", "1.3", "Serialized pipeline caching (classifier.pkl) for sub-millisecond inference"],
        ["Evaluation Plotting", "Matplotlib & Seaborn", "3.10 / 0.13", "Generation of publication-ready Confusion Matrix, ROC, and Loss plots"],
        ["Database Engine", "SQLite 3", "3.x", "Lightweight, zero-config relational database with ACID compliance"]
    ]

    # Create Table
    table = doc.add_table(rows=len(data) + 1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(1.8), Inches(1.5), Inches(0.9), Inches(2.8)]

    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, title_text in enumerate(headers):
        hdr_cells[i].text = title_text
        set_cell_background(hdr_cells[i], "1E3A8A") # Navy Blue
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    # Format Data Rows
    for row_idx, row_data in enumerate(data, start=1):
        row_cells = table.rows[row_idx].cells
        bg_hex = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].text = text
            set_cell_background(row_cells[col_idx], bg_hex)
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=140, right=140)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9.5)
                if col_idx == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(15, 23, 42)
                elif col_idx == 1:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(30, 58, 138)
                else:
                    run.font.color.rgb = RGBColor(51, 65, 85)

    # Set column widths across all rows
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # Set table borders to clean subtle gray
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

    # Add Note footer
    doc.add_paragraph()
    note = doc.add_paragraph()
    note_run = note.add_run("Note: This specification table is prepared for academic evaluation, project documentation, and research conference submissions.")
    note_run.font.name = "Calibri"
    note_run.font.size = Pt(8.5)
    note_run.font.italic = True
    note_run.font.color.rgb = RGBColor(148, 163, 184)

    doc.save(output_path)
    print(f"Document saved successfully at: {output_path}")

if __name__ == "__main__":
    generate_doc()
